from pathlib import Path

from pptx import Presentation

from app.ingestion.document import Document


def load_pptx(file_path: str) -> list[Document]:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"PPTX file not found: {file_path}")

    prs = Presentation(str(path))
    documents: list[Document] = []
    total_slides = len(prs.slides)

    for slide_number, slide in enumerate(prs.slides, start=1):
        text_parts = []

        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                text_parts.append(shape.text_frame.text.strip())

        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        slide_text = "\n".join(text_parts)
        if notes:
            slide_text += f"\n\n[Speaker notes]: {notes}"
        slide_text = slide_text.strip()

        if not slide_text:
            continue

        documents.append(
            Document(
                content=slide_text,
                metadata={
                    "source_type": "pptx",
                    "source_name": path.name,
                    "slide_number": slide_number,
                    "total_slides": total_slides,
                },
            )
        )

    return documents