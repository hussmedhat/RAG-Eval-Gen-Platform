from pathlib import Path

from docx import Document as DocxFile

from app.ingestion.document import Document


def load_docx(file_path: str) -> list[Document]:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"DOCX file not found: {file_path}")

    docx_file = DocxFile(str(path))
    documents: list[Document] = []

    paragraphs = [p.text.strip() for p in docx_file.paragraphs]
    non_empty_paragraphs = [p for p in paragraphs if p]

    for idx, text in enumerate(non_empty_paragraphs, start=1):
        documents.append(
            Document(
                content=text,
                metadata={
                    "source_type": "docx",
                    "source_name": path.name,
                    "paragraph_index": idx,
                    "total_paragraphs": len(non_empty_paragraphs),
                },
            )
        )

    for table_idx, table in enumerate(docx_file.tables, start=1):
        rows_text = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows_text.append(" | ".join(cells))
        table_text = "\n".join(rows_text).strip()

        if table_text:
            documents.append(
                Document(
                    content=table_text,
                    metadata={
                        "source_type": "docx",
                        "source_name": path.name,
                        "table_index": table_idx,
                    },
                )
            )

    return documents